from __future__ import annotations

import hashlib
import io
import json
import os
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS_DIR = Path(os.environ.get("ADPEAK_DOWNLOADS_DIR", r"C:\Users\Lenovo\Downloads"))
INDICADORES_ZIP = Path(os.environ.get(
    "ADPEAK_OFFICIAL_INDICADORES_ZIP",
    DOWNLOADS_DIR / "indicadores-20260628T002121Z-3-001.zip",
))

BACKEND_CATALOG_TARGET = REPO_ROOT / "apps/backend/src/official-catalog.generated.ts"
FRONTEND_CATALOG_TARGET = REPO_ROOT / "apps/frontend/src/catalog/officialCatalog.generated.ts"
BACKEND_DATA_TARGET = REPO_ROOT / "apps/backend/src/official-data.generated.ts"
FRONTEND_DATA_TARGET = REPO_ROOT / "apps/frontend/src/catalog/officialData.generated.ts"
CLASSIFICATION_MATRIX_TARGET = REPO_ROOT / "docs/project/indicator-classification-matrix.md"

CODE_RE = re.compile(r"\b\d+(?:\.\d+){3,}\b")
PLANTEL_RE = re.compile(r"\bBACH(?:ILLERATO)?\s*\.?\s*(\d+)\b|\bBachillerato\s+(\d+)\b", re.I)
PLANTEL_LIST_RE = re.compile(r"\bBachillerato\s+((?:\d+\s*,?\s*){2,})", re.I)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", re.I)
CONTACT_RE = re.compile(r"\b(?:ext\.?|extension|tel(?:efono)?\.?|celular|correo)\b", re.I)
PHONE_NUMBER_RE = re.compile(r"(?<!\d)(?:\+?52\s*)?(?:\d[\s().-]*){8,}\d(?:\.0)?(?!\d)")
MOJIBAKE_MARKERS = ("\u00c3", "\u00c2", "\u00e2", "\ufffd")
KNOWN_TEXT_FIXES = {
    "abadono escolar": "abandono escolar",
    "Abadono escolar": "Abandono escolar",
}
HEADER_TOKENS = {
    "accion",
    "actividad",
    "alumno",
    "avance",
    "bachillerato",
    "cantidad",
    "competencia",
    "correo",
    "descripcion",
    "docente",
    "evidencia",
    "grupo",
    "hombre",
    "indicador",
    "matricula",
    "meta",
    "modalidad",
    "mujer",
    "nombre",
    "observacion",
    "periodo",
    "plantel",
    "programa",
    "responsable",
    "seguimiento",
    "semestre",
    "total",
    "estudiante",
}

READONLY_TOKENS = {
    "delegacion",
    "plantel",
    "programa",
    "semestre",
    "turno",
}

NUMBER_TOKENS = {
    "alumna",
    "alumno",
    "avance",
    "cantidad",
    "hombre",
    "meta",
    "mujer",
    "numero",
    "participante",
    "porcentaje",
    "total",
}

PRIVATE_TOKENS = {
    "correo",
    "curp",
    "cuenta",
    "email",
    "extension",
    "telefonico",
    "telefono",
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").strip()
    text = " ".join(text.split())
    text = repair_mojibake(text)

    for wrong, right in KNOWN_TEXT_FIXES.items():
        text = text.replace(wrong, right)

    return unicodedata.normalize("NFC", text)


def repair_mojibake(value: str) -> str:
    if not any(marker in value for marker in MOJIBAKE_MARKERS):
        return value

    for encoding in ("latin1", "cp1252"):
        try:
            repaired = value.encode(encoding).decode("utf-8")
        except UnicodeError:
            continue

        if sum(repaired.count(marker) for marker in MOJIBAKE_MARKERS) < sum(value.count(marker) for marker in MOJIBAKE_MARKERS):
            return repaired

    return value


def normalize_key(value: str) -> str:
    text = unicodedata.normalize("NFD", value.lower())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def json_ts(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2)


def write_text(path: Path, text: str) -> None:
    if any(marker in text for marker in MOJIBAKE_MARKERS):
        raise RuntimeError(f"Refusing to write mojibake into {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")


def markdown_cell(value: Any) -> str:
    return clean_text(value).replace("|", "\\|")


def generate_classification_matrix(rows: list[dict[str, Any]], stats: dict[str, Any]) -> str:
    lines = [
        "# Matriz de clasificación de indicadores oficiales",
        "",
        f"Fuente: `{INDICADORES_ZIP.name}`",
        "",
        "Este archivo se regenera mediante `python tools/import-official-data.py`; no se edita manualmente.",
        "",
        "| Código visible | Código fuente | Clasificación | Visible | Responsable | Fuente |",
        "| --- | --- | --- | --- | --- | --- |",
    ]

    for row in rows:
        visible_code = row.get("code") if row.get("visible") else "-"
        lines.append("| " + " | ".join([
            markdown_cell(visible_code),
            markdown_cell(row.get("sourceCode")),
            markdown_cell(row.get("classification")),
            "sí" if row.get("visible") else "no",
            markdown_cell(row.get("responsible")),
            markdown_cell(row.get("activity")),
        ]) + " |")

    lines.extend([
        "",
        "## Resumen",
        "",
        f"- Filas físicas trazables: `{stats['sourceRows']}`.",
        f"- Indicadores operativos visibles y únicos: `{stats['uniqueIndicators']}`.",
        f"- Fuentes operativas: `{stats['operationalRows']}`.",
        f"- Plantillas internas: `{stats['templateRows']}`.",
        f"- Variantes de plantilla internas: `{stats['templateVariantRows']}`.",
        f"- Fuentes pendientes de mapeo: `{stats['pendingMappingRows']}`.",
        "- `operational`: aparece en el flujo normal cuando el rol y alcance lo permiten.",
        "- `template` y `template_variant`: estructura interna; nunca aparecen como indicadores independientes.",
        "- `pending_mapping`: fuente conservada sin inventar un indicador o columnas operativas.",
    ])
    return "\n".join(lines)


def catalog_stats(rows: list[dict[str, Any]]) -> dict[str, Any]:
    operational_rows = [row for row in rows if row.get("classification") == "operational"]
    contributors = sorted(
        {
            person.strip()
            for row in operational_rows
            for person in row["contributors"].split(",")
            if person.strip()
        },
        key=lambda item: item.casefold(),
    )

    return {
        "sourceRows": len(rows),
        "uniqueRows": len({row["dedupeKey"] for row in rows}),
        "duplicateRows": sum(1 for row in rows if row["isDuplicate"]),
        "uniqueSourceCodes": len({row["sourceCode"] for row in rows if row.get("sourceCode")}),
        "uniqueIndicators": len({row["code"] for row in operational_rows if row["code"]}),
        "operationalRows": len(operational_rows),
        "templateRows": sum(1 for row in rows if row.get("classification") == "template"),
        "templateVariantRows": sum(1 for row in rows if row.get("classification") == "template_variant"),
        "pendingMappingRows": sum(1 for row in rows if row.get("classification") == "pending_mapping"),
        "uniqueResponsibles": len({row["responsible"] for row in operational_rows if row["responsible"]}),
        "uniqueContributors": len(contributors),
        "uniqueActivities": len({row["activity"] for row in operational_rows if row["activity"]}),
        "blankActivities": sum(1 for row in rows if not row["activity"]),
        "workbookOnlyIndicators": sum(1 for row in rows if "workbook_only_indicator" in row["dataQuality"]),
    }


def frontend_catalog_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    operational_rows = [
        row
        for row in rows
        if row.get("classification") == "operational"
        and row.get("visible") is True
        and not is_synthetic_indicator_code(row.get("code", ""))
    ]
    aliases = person_aliases(operational_rows)
    sanitized: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, str]] = set()

    for row in operational_rows:
        responsible = aliases.get(row["responsible"], "Responsable sin asignar")
        contributors = ", ".join(
            aliases.get(person, "Responsable sin asignar")
            for person in split_people(row["contributors"])
        )
        dedupe_key = (row["code"], responsible, contributors, row["activity"])

        if dedupe_key in seen:
            continue

        seen.add(dedupe_key)
        sanitized.append({
            "code": row["code"],
            "name": row["name"],
            "responsible": responsible,
            "contributors": contributors,
            "activity": row["activity"],
        })

    return sanitized


def is_synthetic_indicator_code(code: str) -> bool:
    return code.startswith("FMT-") or "-FMT-" in code or code.startswith("TMP-")


def inferred_indicator_code_from_table(sheets: list[dict[str, Any]]) -> str:
    """Return one code only when the workbook table identifies it unambiguously."""
    candidates: set[str] = set()

    for sheet in sheets:
        if not sheet.get("table"):
            continue

        descriptions = [clean_text(value) for value in sheet.get("codeDescriptions", [])]
        for code in sheet.get("_codes", []):
            if any(code in description for description in descriptions):
                candidates.add(code)

    return next(iter(candidates)) if len(candidates) == 1 else ""


def person_aliases(rows: list[dict[str, Any]]) -> dict[str, str]:
    people = sorted(
        {
            person
            for row in rows
            for person in [row["responsible"], *split_people(row["contributors"])]
            if person
        },
        key=lambda value: normalize_key(value),
    )
    return {person: f"Responsable {index:02d}" for index, person in enumerate(people, start=1)}


def split_people(value: str) -> list[str]:
    return [clean_text(part) for part in value.split(",") if clean_text(part)]


def catalog_plantel_scopes(rows: list[dict[str, Any]], detected_scopes: dict[str, list[int]]) -> dict[str, list[int]]:
    catalog_codes = {row["code"] for row in rows if row["code"] and row.get("classification") == "operational"}
    scopes: dict[str, set[int]] = {}

    for code, plantel_ids in detected_scopes.items():
        if code in catalog_codes:
            scopes.setdefault(code, set()).update(plantel_ids)

    return {
        code: sorted(plantel_ids)
        for code, plantel_ids in sorted(scopes.items(), key=lambda item: normalize_key(item[0]))
    }


def is_format_source(source_path: str, source_label: str = "", activity: str = "") -> bool:
    normalized_path = unicodedata.normalize("NFC", source_path)
    parts = [normalize_key(part) for part in normalized_path.split("/") if part]
    stem = normalize_key(Path(normalized_path).stem)
    label = normalize_key(source_label)
    activity_key = normalize_key(activity)

    if "formatos" in parts:
        return True
    if "formato" in stem.split() or stem.startswith("formato "):
        return True
    if "formato" in label.split() or label.startswith("formato "):
        return True
    if "opcion de llenado" in stem or "opcion de llenado" in label or "opcion de llenado" in activity_key:
        return True
    return False


def classify_catalog_source(source_code: str, official_code: str, source_path: str, source_label: str, activity: str) -> tuple[str, list[str]]:
    reasons: list[str] = []
    if source_code.startswith("FMT-"):
        reasons.append("synthetic_code")
    if "-FMT-" in source_code:
        reasons.append("variant_source_code")
    if is_format_source(source_path, source_label, activity):
        reasons.append("format_source")
    if not official_code:
        reasons.append("no_official_code")
        return "pending_mapping", reasons
    if "format_source" in reasons:
        return ("template_variant" if source_code != official_code or "-FMT-" in source_code else "template"), reasons
    if source_code != official_code:
        reasons.append("merged_to_official_code")
    return "operational", reasons


def plantel_id_from_name(name: str) -> int | None:
    normalized = normalize_key(name)

    if "linea" in normalized:
        return 36
    if "iuba" in normalized or "bellas artes" in normalized:
        return 37

    match = re.search(r"\b(\d+)\b", normalized)
    if not match:
        return None

    number = int(match.group(1))
    legacy_ids = {16: 1, 4: 2, 1: 3, 33: 4}
    if number in legacy_ids:
        return legacy_ids[number]
    if not 1 <= number <= 35:
        return None

    generated_numbers = [candidate for candidate in range(1, 36) if candidate not in legacy_ids]
    return 5 + generated_numbers.index(number)


RESPONSIBLE_FOLDER_ALIASES = {
    "adriana ruiz": "Adriana Ruiz Rivera",
    "ariadna zuniga": "Ariadna Zúñiga Torres",
    "angel ordonez": "Angel Ordoñez",
    "daniela navarro": "Daniela Nohemi Navarro Castillo",
    "liliana rojas": "Liliana Yunuen Rojas Maciel",
    "oscar pedraza": "Oscar Pedraza Farías",
    "carlos nava": "Carlos Hernández Nava",
    "oscar mendoza": "Oscar Gustavo Mendoza Barajas",
    "laura calvario solo responsable": "Laura Gabriela Calvario",
    "marcial avina": "Marcial Aviña Iglesias",
    "oscar delgado": "Oscar Delgado Sánchez",
}


def workbook_catalog_rows(workbook_templates: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    for source_row, (source_code, template) in enumerate(
        sorted(workbook_templates.items(), key=lambda item: normalize_key(item[0])),
        start=2,
    ):
        source_path = clean_text(template.get("sourcePath"))
        if not source_path.startswith("Indicadores/"):
            continue

        responsible = responsible_from_source_path(source_path)
        source_label = clean_text(template.get("sourceLabel"))
        official_code = clean_text(template.get("officialCode"))
        activity = (
            clean_text(template.get("activityLabel"))
            or source_label
            or clean_text(Path(source_path).name)
            or "Formato oficial importado"
        )
        classification, classification_reason = classify_catalog_source(
            source_code,
            official_code,
            source_path,
            source_label,
            activity,
        )
        code = official_code if classification == "operational" and official_code else source_code
        name = clean_text(template.get("indicatorName")) or official_code or source_code
        dedupe_key = "\u241f".join((source_code, code, name, responsible, "Planteles", activity, classification))
        rows.append({
            "sourceRow": source_row,
            "code": code,
            "sourceCode": source_code,
            "officialCode": official_code or None,
            "name": name,
            "responsible": responsible,
            "contributors": "Planteles",
            "activity": activity,
            "classification": classification,
            "visible": classification == "operational",
            "classificationReason": classification_reason,
            "dedupeKey": hashlib.sha256(dedupe_key.encode("utf-8")).hexdigest()[:16],
            "isDuplicate": False,
            "duplicateOfSourceRow": None,
            "dataQuality": sorted(set(list(template.get("quality") or []) + ["source_zip_indicator"])),
        })

    return rows


def canonicalize_workbook_templates(workbook_templates: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    canonical = dict(workbook_templates)
    by_official_code: dict[str, list[tuple[str, dict[str, Any], str]]] = defaultdict(list)

    for source_code, template in workbook_templates.items():
        official_code = clean_text(template.get("officialCode"))
        if not official_code:
            continue
        source_path = clean_text(template.get("sourcePath"))
        source_label = clean_text(template.get("sourceLabel"))
        activity = clean_text(template.get("activityLabel")) or source_label or clean_text(Path(source_path).name)
        classification, _reason = classify_catalog_source(source_code, official_code, source_path, source_label, activity)
        by_official_code[official_code].append((source_code, template, classification))

    for official_code, templates in by_official_code.items():
        operational_templates = [
            (source_code, template)
            for source_code, template, classification in templates
            if classification == "operational"
        ]
        if not operational_templates:
            continue

        selected_code, selected_template = max(
            operational_templates,
            key=lambda item: (
                item[0] == official_code,
                len(item[1].get("initialRows") or []),
                len(item[1].get("columns") or []),
            ),
        )
        if selected_code == official_code:
            continue

        current_template = canonical.get(official_code)
        if current_template is not None:
            source_path = clean_text(current_template.get("sourcePath")) or official_code
            hidden_suffix = hashlib.sha256(source_path.encode("utf-8")).hexdigest()[:8].upper()
            hidden_key = f"{official_code}-FMT-{hidden_suffix}"
            while hidden_key in canonical and hidden_key != selected_code:
                hidden_suffix = hashlib.sha256((source_path + hidden_key).encode("utf-8")).hexdigest()[:8].upper()
                hidden_key = f"{official_code}-FMT-{hidden_suffix}"
            hidden_template = dict(current_template)
            hidden_template["indicatorCode"] = hidden_key
            hidden_template["officialCode"] = official_code
            hidden_template["quality"] = sorted(set(list(hidden_template.get("quality") or []) + ["canonical_format_hidden"]))
            canonical[hidden_key] = hidden_template

        promoted_template = dict(selected_template)
        promoted_template["indicatorCode"] = official_code
        promoted_template["officialCode"] = official_code
        promoted_template["quality"] = sorted(set(list(promoted_template.get("quality") or []) + ["canonical_operational_template"]))
        canonical[official_code] = promoted_template

        if selected_code in canonical and selected_code != official_code:
            del canonical[selected_code]

    return canonical


def responsible_from_source_path(source_path: str) -> str:
    parts = [part for part in unicodedata.normalize("NFC", source_path).split("/") if part]
    folder = parts[1] if len(parts) > 2 and parts[0] == "Indicadores" else parts[0] if parts else ""
    normalized = normalize_key(folder)
    return RESPONSIBLE_FOLDER_ALIASES.get(normalized, folder.replace(" - solo responsable", "").strip() or "Pendiente de asignar")


def workbook_summaries() -> tuple[
    list[dict[str, Any]],
    dict[str, list[int]],
    list[dict[str, Any]],
    list[dict[str, Any]],
    dict[str, dict[str, Any]],
]:
    summaries: list[dict[str, Any]] = []
    evidence_groups: dict[str, dict[str, Any]] = {}
    scopes: dict[str, set[int]] = defaultdict(set)
    template_candidates: list[dict[str, Any]] = []
    workbook_templates: dict[str, dict[str, Any]] = {}

    archives = official_workbook_archives()
    if not archives:
        return summaries, {}, [], [], {}

    for archive_path, source_scope, source_prefix in archives:
        with zipfile.ZipFile(archive_path) as archive:
            infos = [info for info in archive.infolist() if not info.is_dir()]
            for info in infos:
                category = category_from_path(info.filename, source_prefix)
                group = evidence_groups.setdefault(
                    category,
                    {"category": category, "fileCount": 0, "totalBytes": 0, "byExtension": Counter()},
                )
                group["fileCount"] += 1
                group["totalBytes"] += info.file_size
                group["byExtension"][Path(info.filename).suffix.lower() or "<none>"] += 1

            for info in infos:
                if Path(info.filename).suffix.lower() != ".xlsx":
                    continue
                payload = archive.read(info)
                digest = hashlib.sha256(payload).hexdigest()
                workbook = load_workbook(io.BytesIO(payload), data_only=False, read_only=False)
                reference_codes: set[str] = set()
                indicator_codes: set[str] = set()
                planteles: set[str] = set()
                sheets = []
                formulas = 0

                for sheet in workbook.worksheets:
                    sheet_summary = summarize_sheet(sheet)
                    reliable_sheet_codes = sheet_summary.pop("indicatorCodes")
                    sheet_codes = sheet_summary.pop("codes")
                    sheet_summary["_indicatorCodes"] = sorted(reliable_sheet_codes, key=normalize_key)
                    sheet_summary["_codes"] = sorted(sheet_codes, key=normalize_key)
                    indicator_codes.update(reliable_sheet_codes)
                    reference_codes.update(sheet_codes)
                    planteles.update(sheet_summary.pop("planteles"))
                    formulas += sheet_summary["formulaCells"]
                    sheets.append(sheet_summary)

                filename_codes = CODE_RE.findall(unicodedata.normalize("NFC", info.filename))
                if filename_codes:
                    indicator_codes.update(filename_codes)

                if not planteles and source_scope:
                    planteles.add(source_scope)
                table_sheets = [sheet for sheet in sheets if sheet.get("table")]
                if not indicator_codes:
                    inferred_code = inferred_indicator_code_from_table(sheets)
                    if inferred_code:
                        indicator_codes.add(inferred_code)
                source_codes = sorted(indicator_codes, key=normalize_key)
                template_codes = source_codes or (
                    [synthetic_workbook_code(info.filename, digest)] if table_sheets else []
                )

                for source_index, source_code in enumerate(template_codes, start=1):
                    original_code = source_code if source_code in source_codes else ""
                    template_code = unique_workbook_template_code(
                        source_code,
                        info.filename,
                        digest,
                        workbook_templates,
                        original_code=original_code,
                        source_index=source_index,
                    )

                    detected_plantel_ids = [
                        plantel_id
                        for plantel_id in (plantel_id_from_name(plantel) for plantel in planteles)
                        if plantel_id is not None
                    ]
                    if source_scope:
                        source_scope_id = plantel_id_from_name(source_scope)
                        if source_scope_id is not None:
                            detected_plantel_ids.append(source_scope_id)

                    for plantel_id in detected_plantel_ids:
                        scopes[template_code].add(plantel_id)
                        if original_code:
                            scopes[original_code].add(plantel_id)

                    template_candidates.append(
                        {
                            "indicatorCode": template_code,
                            "officialCode": original_code or None,
                            "sourcePath": source_path_for_summary(info.filename, source_prefix),
                            "planteles": sorted(planteles, key=str.casefold),
                            "classification": classify_workbook(sheets),
                            "confidence": "detected-code" if original_code else "source-table-pending-code",
                            "headerRows": first_header_rows(sheets),
                        }
                    )
                    table_template = template_from_workbook_sheets(
                        template_code,
                        sheets,
                        source_path_for_summary(info.filename, source_prefix),
                        official_code=original_code,
                        source_index=source_index,
                    )
                    if table_template and template_code not in workbook_templates:
                        workbook_templates[template_code] = table_template

                summaries.append(
                    {
                        "id": digest[:16],
                        "sourceLabel": unicodedata.normalize("NFC", Path(info.filename).name),
                        "sourcePathHash": digest,
                        "category": category_from_path(info.filename, source_prefix),
                        "sizeBytes": info.file_size,
                        "detectedIndicatorCodes": sorted(indicator_codes),
                        "detectedReferenceCodes": sorted(reference_codes - indicator_codes),
                        "detectedPlanteles": sorted(planteles, key=str.casefold),
                        "formulaCells": formulas,
                        "sheets": [
                            {
                                key: value
                                for key, value in sheet.items()
                                if not key.startswith("_")
                            }
                            for sheet in sheets
                        ],
                        "privacy": "Row-level personal data remains in private storage only.",
                    }
                )

    groups = []
    for value in evidence_groups.values():
        by_ext = dict(sorted(value["byExtension"].items()))
        groups.append(
            {
                "category": value["category"],
                "fileCount": value["fileCount"],
                "totalBytes": value["totalBytes"],
                "byExtension": by_ext,
                "sampleFileTypes": sorted(by_ext)[:6],
            }
        )

    return (
        summaries,
        {code: sorted(ids) for code, ids in scopes.items()},
        sorted(groups, key=lambda item: item["category"].casefold()),
        template_candidates,
        workbook_templates,
    )


def unique_workbook_template_code(
    source_code: str,
    source_path: str,
    digest: str,
    existing_templates: dict[str, dict[str, Any]],
    *,
    original_code: str = "",
    source_index: int = 1,
) -> str:
    if source_code not in existing_templates:
        return source_code

    candidate = synthetic_workbook_code(source_path, digest, original_code or source_code, source_index)
    counter = 2
    while candidate in existing_templates:
        candidate = synthetic_workbook_code(source_path, digest, original_code or source_code, source_index + counter)
        counter += 1
    return candidate


def official_workbook_archives() -> list[tuple[Path, str | None, str]]:
    archives: list[tuple[Path, str | None, str]] = []
    if INDICADORES_ZIP.exists():
        archives.append((INDICADORES_ZIP, None, "Indicadores"))
    return archives


def source_path_for_summary(filename: str, source_prefix: str) -> str:
    clean = unicodedata.normalize("NFC", repair_mojibake(filename))
    if source_prefix == "Indicadores" and clean.startswith("indicadores/"):
        clean = clean.removeprefix("indicadores/")
    return clean if clean.startswith(f"{source_prefix}/") else f"{source_prefix}/{clean}"


def summarize_sheet(sheet: Any) -> dict[str, Any]:
    codes: set[str] = set()
    planteles: set[str] = set()
    rows_seen = 0
    numeric_cells = 0
    text_cells = 0
    formula_cells = 0
    header_rows: list[list[str]] = []
    non_empty_rows: list[dict[str, Any]] = []
    code_descriptions: list[str] = []
    activity_descriptions: list[str] = []
    columns_observed = 0
    indicator_codes: set[str] = set()

    max_row = min(sheet.max_row or 80, 120)
    max_col = min(sheet.max_column or 30, 30)
    merged_ranges = merged_ranges_for_sheet(sheet, max_row, max_col)

    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=False), start=1):
        values = [clean_text(cell.value) for cell in row]
        non_empty = [value for value in values if value]
        if not non_empty:
            continue
        rows_seen += 1
        for column_index, value in enumerate(values, start=1):
            if value:
                columns_observed = max(columns_observed, column_index)
        row_text = " ".join(non_empty)
        non_empty_rows.append({"row": row_number, "values": trim_trailing_blanks(values)})
        row_codes = CODE_RE.findall(row_text)
        codes.update(row_codes)
        if row_codes and "indicador" in normalize_key(row_text):
            indicator_codes.update(row_codes)
        if row_codes and not is_reference_note_row(normalize_key(row_text)) and len(code_descriptions) < 5:
            code_descriptions.append(row_text)
        activity_description = activity_description_from_row(row_text)
        if activity_description and len(activity_descriptions) < 3:
            activity_descriptions.append(activity_description)
        for match in PLANTEL_RE.finditer(row_text):
            number = match.group(1) or match.group(2)
            if number:
                planteles.add(f"Bachillerato {int(number)}")
        for match in PLANTEL_LIST_RE.finditer(row_text):
            for number in re.findall(r"\d+", match.group(1)):
                planteles.add(f"Bachillerato {int(number)}")
        normalized_row_text = normalize_key(row_text)
        if "bachillerato" in normalized_row_text and "1,2" in row_text.replace(" ", "") and re.search(r"\b35\b", row_text):
            for number in range(1, 36):
                planteles.add(f"Bachillerato {number}")
        if "iuba" in normalized_row_text or "bellas artes" in normalized_row_text:
            planteles.add("IUBA Bachillerato")
        if "bachillerato en linea" in normalized_row_text:
            planteles.add("Bachillerato en línea")

        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_cells += 1
            elif isinstance(cell.value, (int, float)):
                numeric_cells += 1
            elif cell.value is not None:
                text_cells += 1

        candidate = sanitize_header_row(non_empty)
        if candidate and len(header_rows) < 6:
            header_rows.append(candidate)

    return {
        "name": clean_text(sheet.title),
        "nonEmptyRows": rows_seen,
        "columnsObserved": columns_observed,
        "sampleHeaders": header_rows[0] if header_rows else [],
        "headerRows": header_rows,
        "numericCells": numeric_cells,
        "textCells": text_cells,
        "formulaCells": formula_cells,
        "table": extract_table_from_rows(non_empty_rows, merged_ranges),
        "codeDescriptions": code_descriptions,
        "activityDescriptions": activity_descriptions,
        "indicatorCodes": indicator_codes,
        "codes": codes,
        "planteles": planteles,
    }


def merged_ranges_for_sheet(sheet: Any, max_row: int, max_col: int) -> list[dict[str, Any]]:
    ranges: list[dict[str, Any]] = []

    for merged in sheet.merged_cells.ranges:
        min_row, min_col, max_merged_row, max_merged_col = (
            merged.min_row,
            merged.min_col,
            merged.max_row,
            merged.max_col,
        )

        if min_row > max_row or min_col > max_col:
            continue

        label = clean_text(sheet.cell(min_row, min_col).value)
        if not label:
            continue

        ranges.append({
            "min_row": min_row,
            "max_row": min(max_merged_row, max_row),
            "min_col": min_col,
            "max_col": min(max_merged_col, max_col),
            "label": label,
        })

    return ranges


def activity_description_from_row(row_text: str) -> str:
    match = re.search(r"(?i)\bACTIVIDADES?\s*:\s*(.+)", row_text)
    if not match:
        return ""

    activity = match.group(1)
    activity = re.split(r"(?i)\b(registre|nota|indicador|c[oó]digo)\b", activity)[0]
    activity = clean_text(activity).strip(" .:-")
    return activity[:180]


def trim_trailing_blanks(values: list[str]) -> list[str]:
    trimmed = list(values)
    while trimmed and not trimmed[-1]:
        trimmed.pop()
    return trimmed


def fallback_header_label(_index: int) -> str:
    return "Registro"


def strip_empty_header_rows(header_rows: list[list[dict[str, Any]]]) -> list[list[dict[str, Any]]]:
    non_empty_rows = [
        row
        for row in header_rows
        if any(clean_text(cell.get("label", "")) for cell in row)
    ]
    normalized_rows: list[list[dict[str, Any]]] = []

    for row_index, row in enumerate(non_empty_rows):
        remaining_rows = len(non_empty_rows) - row_index
        normalized_row: list[dict[str, Any]] = []
        for source_cell in row:
            cell = dict(source_cell)
            if "rowspan" in cell:
                rowspan = min(max(1, int(cell.get("rowspan", 1) or 1)), remaining_rows)
                if rowspan > 1:
                    cell["rowspan"] = rowspan
                else:
                    cell.pop("rowspan", None)
            normalized_row.append(cell)
        normalized_rows.append(normalized_row)

    return normalized_rows


def header_layout_is_valid(
    header_rows: list[list[dict[str, Any]]],
    column_count: int,
) -> bool:
    """Validate that every header row resolves to exactly the table width."""
    if not header_rows or column_count <= 0:
        return False

    occupied_until = [0] * column_count

    for row_index, row in enumerate(header_rows):
        cursor = 0
        covered: set[int] = {
            column_index
            for column_index, until in enumerate(occupied_until)
            if until > row_index
        }

        for cell in row:
            while cursor < column_count and cursor in covered:
                cursor += 1

            colspan = max(1, int(cell.get("colspan", 1) or 1))
            rowspan = max(1, int(cell.get("rowspan", 1) or 1))
            end = cursor + colspan

            if cursor >= column_count or end > column_count:
                return False
            if any(column_index in covered for column_index in range(cursor, end)):
                return False

            for column_index in range(cursor, end):
                covered.add(column_index)
                if rowspan > 1:
                    occupied_until[column_index] = max(
                        occupied_until[column_index],
                        row_index + rowspan,
                    )
            cursor = end

        if len(covered) != column_count:
            return False

    return True


def repair_header_geometry(
    header_rows: list[list[dict[str, Any]]],
    column_count: int,
) -> list[list[dict[str, Any]]]:
    """Repair Excel group cells that incorrectly overlap their subheaders."""
    normalized = strip_empty_header_rows(header_rows)
    if header_layout_is_valid(normalized, column_count):
        return normalized

    repaired: list[list[dict[str, Any]]] = []
    for row in normalized:
        repaired_row: list[dict[str, Any]] = []
        for source_cell in row:
            cell = dict(source_cell)
            if int(cell.get("colspan", 1) or 1) > 1 and int(cell.get("rowspan", 1) or 1) > 1:
                cell.pop("rowspan", None)
            repaired_row.append(cell)
        repaired.append(repaired_row)

    if header_layout_is_valid(repaired, column_count):
        return repaired

    # An invalid grouped header is worse than the safe column-label fallback.
    return []


def project_header_rows(
    header_rows: list[list[dict[str, Any]]],
    source_width: int,
    kept_source_indices: list[int],
) -> list[list[dict[str, Any]]]:
    """Remove hidden columns from structured headers without losing group labels."""
    if not header_rows or source_width <= 0:
        return strip_empty_header_rows(header_rows)

    occupied_until = [0] * source_width
    projected_rows: list[list[dict[str, Any]]] = []

    for row_index, row in enumerate(header_rows):
        cursor = 0
        projected_cells: list[tuple[int, dict[str, Any]]] = []

        for source_cell in row:
            while cursor < source_width and occupied_until[cursor] > row_index:
                cursor += 1
            if cursor >= source_width:
                break

            colspan = max(1, int(source_cell.get("colspan", 1) or 1))
            rowspan = max(1, int(source_cell.get("rowspan", 1) or 1))
            end = min(source_width, cursor + colspan)
            covered_source_indices = set(range(cursor, end))
            projected_positions = [
                projected_index
                for projected_index, source_index in enumerate(kept_source_indices)
                if source_index in covered_source_indices
            ]

            if projected_positions:
                projected_cell = {
                    key: value
                    for key, value in source_cell.items()
                    if key not in {"colspan", "rowspan"}
                }
                if len(projected_positions) > 1:
                    projected_cell["colspan"] = len(projected_positions)
                if rowspan > 1:
                    projected_cell["rowspan"] = rowspan
                projected_cells.append((min(projected_positions), projected_cell))

            if rowspan > 1:
                for source_index in range(cursor, end):
                    occupied_until[source_index] = max(
                        occupied_until[source_index],
                        row_index + rowspan,
                    )
            cursor = end

        projected_rows.append([
            cell
            for _, cell in sorted(projected_cells, key=lambda item: item[0])
        ])

    return strip_empty_header_rows(projected_rows)


def values_for_column(rows: list[dict[str, Any]], key: str) -> list[str]:
    return [clean_text(row.get(key, "")) for row in rows]


def all_blank(values: list[str]) -> bool:
    return all(not value for value in values)


def is_blank_readonly_context_column(column: dict[str, Any], rows: list[dict[str, Any]]) -> bool:
    if column.get("type") != "readonly":
        return False

    normalized = normalize_key(f"{column.get('label', '')} {column.get('key', '')}")
    if "delegacion" not in normalized:
        return False

    return all_blank(values_for_column(rows, column["key"]))


def readable_label_from_key(key: str) -> str:
    label = re.sub(r"_\d+$", "", key).replace("_", " ").strip()
    if not label:
        return "Registro"
    replacements = {
        "ano": "año",
        "matricula": "matrícula",
        "titulacion": "titulación",
        "descripcion": "descripción",
        "capacitacion": "capacitación",
    }
    words = []
    for word in label.split():
        if word in {"m", "h", "t"}:
            words.append(word.upper())
        else:
            words.append(replacements.get(word, word).capitalize())
    return " ".join(words)


def disambiguate_repeated_labels(columns: list[dict[str, Any]]) -> list[dict[str, Any]]:
    counts = Counter(normalize_key(column["label"]) for column in columns)
    next_columns: list[dict[str, Any]] = []

    for column in columns:
        normalized = normalize_key(column["label"])
        if counts[normalized] <= 1:
            next_columns.append(column)
            continue

        replacement_label = readable_label_from_key(column["key"])
        if normalize_key(column["label"]) not in normalize_key(replacement_label):
            replacement_label = f"{replacement_label} {column['label']}"

        next_columns.append({
            **column,
            "label": replacement_label,
        })

    return next_columns


def normalize_extracted_table(
    columns: list[dict[str, Any]],
    initial_rows: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[list[dict[str, Any]]]]:
    rows = [dict(row) for row in initial_rows]
    kept_columns: list[dict[str, Any]] = []
    kept_source_indices: list[int] = []
    kept_by_signature: dict[tuple[str, str], int] = {}
    dropped_columns: set[str] = set()

    for source_index, column in enumerate(columns):
        next_column = dict(column)
        normalized_label = normalize_key(next_column["label"])
        base_key = re.sub(r"_\d+$", "", next_column["key"])
        signature = (normalized_label, base_key)
        can_dedupe = next_column.get("type") in {"readonly", "text"}

        if is_blank_readonly_context_column(next_column, rows):
            dropped_columns.add(next_column["key"])
            continue

        existing_index = kept_by_signature.get(signature) if can_dedupe else None
        if existing_index is not None:
            existing_column = kept_columns[existing_index]
            existing_values = values_for_column(rows, existing_column["key"])
            candidate_values = values_for_column(rows, next_column["key"])

            if all_blank(candidate_values) or candidate_values == existing_values:
                dropped_columns.add(next_column["key"])
                continue

            if all_blank(existing_values):
                dropped_columns.add(existing_column["key"])
                kept_columns[existing_index] = next_column
                kept_source_indices[existing_index] = source_index
                continue

        kept_by_signature[signature] = len(kept_columns)
        kept_columns.append(next_column)
        kept_source_indices.append(source_index)

    if dropped_columns:
        for row in rows:
            for key in dropped_columns:
                row.pop(key, None)
        header_rows = project_header_rows(header_rows, len(columns), kept_source_indices)
        kept_columns = disambiguate_repeated_labels(kept_columns)
    else:
        header_rows = strip_empty_header_rows(header_rows)

    header_rows = repair_header_geometry(header_rows, len(kept_columns))

    return kept_columns, rows, header_rows


def extract_table_from_rows(rows: list[dict[str, Any]], merged_ranges: list[dict[str, Any]] | None = None) -> dict[str, Any] | None:
    if not rows:
        return None

    best_index = -1
    best_score = 0
    for index, row in enumerate(rows):
        values = row["values"]
        normalized = " ".join(normalize_key(value) for value in values if value)
        token_score = sum(1 for token in HEADER_TOKENS if token in normalized)
        core_score = sum(1 for token in ("plantel", "actividad", "meta", "avance", "observacion") if token in normalized)
        width_score = min(4, len([value for value in values if value]))
        next_values = adjacent_row_values(rows, index, 1)
        third_values = adjacent_row_values(rows, index, 2)
        merge_bonus = 0
        if next_values and third_values and should_merge_three_header_rows(values, next_values, third_values):
            merge_bonus = 18
        elif next_values and should_merge_header_rows(values, next_values):
            merge_bonus = 12
        score = token_score * 3 + core_score * 4 + width_score + merge_bonus
        if score > best_score and token_score >= 1 and width_score >= 2:
            best_index = index
            best_score = score

    if best_index < 0:
        return None

    header_matrix, header_row_numbers, data_start = combined_header_matrix(rows, best_index)
    header_values = combined_header_labels(header_matrix, header_row_numbers, merged_ranges or [])
    header_rows = structured_header_rows(header_matrix, header_row_numbers, merged_ranges or [])
    column_labels = display_column_labels(header_matrix, header_row_numbers, merged_ranges or [])
    width = len(header_values)
    columns = []
    seen_keys: set[str] = set()
    for index, label in enumerate(header_values):
        if not label:
            label = fallback_header_label(index)
        display_label = column_labels[index] if index < len(column_labels) and column_labels[index] else label
        if normalize_key(display_label).startswith("columna"):
            display_label = fallback_header_label(index)
        columns.append(column_from_label(display_label, seen_keys, key_label=label))
    apply_official_calculated_columns(columns)

    initial_rows = []
    for row in rows[data_start:]:
        values = row["values"][:width]
        if not any(values):
            continue
        normalized = normalize_key(" ".join(values))
        non_empty_values = [value for value in values if value]
        if is_non_capture_row(normalized, non_empty_values):
            continue

        row_object: dict[str, Any] = {}
        for column, value in zip(columns, values + [""] * (width - len(values))):
            row_object[column["key"]] = value_for_column(column, value)
        if any(value not in ("", None) for value in row_object.values()):
            initial_rows.append(row_object)
        if len(initial_rows) >= 60:
            break

    if not initial_rows:
        initial_rows = [{column["key"]: "" for column in columns}]

    columns, initial_rows, header_rows = normalize_extracted_table(columns, initial_rows, header_rows)

    return {
        "headerRow": rows[best_index]["row"],
        "headerRows": header_rows,
        "columns": columns,
        "initialRows": initial_rows,
    }


def is_non_capture_row(normalized: str, non_empty_values: list[str]) -> bool:
    if not non_empty_values:
        return True
    if ("nota" in normalized or "total" in normalized or "totales" in normalized) and len(non_empty_values) <= 2:
        return True
    if is_reference_note_row(normalized):
        return True
    return False


def is_reference_note_row(normalized: str) -> bool:
    return any(
        marker in normalized
        for marker in (
            "la tabla anterior incide",
            "linea de accion",
            "lineas de accion",
            "informe de labores",
            "instrucciones para el llenado",
        )
    )


def combined_header_values(rows: list[dict[str, Any]], best_index: int) -> tuple[list[str], int]:
    header_matrix, header_row_numbers, data_start = combined_header_matrix(rows, best_index)
    return combined_header_labels(header_matrix, header_row_numbers, []), data_start


def combined_header_matrix(rows: list[dict[str, Any]], best_index: int) -> tuple[list[list[str]], list[int], int]:
    primary = rows[best_index]["values"]
    secondary = adjacent_row_values(rows, best_index, 1)
    tertiary = adjacent_row_values(rows, best_index, 2)

    if secondary and tertiary and should_merge_three_header_rows(primary, secondary, tertiary):
        return [primary, secondary, tertiary], [rows[best_index + offset]["row"] for offset in range(3)], best_index + 3

    if not should_merge_header_rows(primary, secondary):
        return [primary], [rows[best_index]["row"]], best_index + 1

    return [primary, secondary], [rows[best_index]["row"], rows[best_index + 1]["row"]], best_index + 2


def combined_header_labels(
    header_rows: list[list[str]],
    header_row_numbers: list[int],
    merged_ranges: list[dict[str, Any]],
) -> list[str]:
    if not merged_ranges:
        return merge_header_rows(header_rows)

    header_ranges = header_ranges_with_vertical_singletons(header_rows, header_row_numbers, merged_ranges)
    grid = expanded_header_grid(header_rows, header_row_numbers, header_ranges)
    labels: list[str] = []

    if not grid:
        return []

    width = max(len(row) for row in grid)
    for col_index in range(width):
        parts: list[str] = []
        normalized_parts: set[str] = set()
        for row in grid:
            label = clean_text(row[col_index] if col_index < len(row) else "")
            normalized = normalize_key(label)
            if not label or normalized in normalized_parts:
                continue
            if len(normalized) > 2 and any(normalized in existing for existing in normalized_parts):
                continue
            parts.append(label)
            normalized_parts.add(normalized)
        labels.append(" ".join(parts) or fallback_header_label(col_index))

    return labels


def display_column_labels(
    header_rows: list[list[str]],
    header_row_numbers: list[int],
    merged_ranges: list[dict[str, Any]],
) -> list[str]:
    if not merged_ranges:
        return merge_header_rows(header_rows)

    header_ranges = header_ranges_with_vertical_singletons(header_rows, header_row_numbers, merged_ranges)
    grid = expanded_header_grid(header_rows, header_row_numbers, header_ranges)
    labels: list[str] = []

    if not grid:
        return []

    width = max(len(row) for row in grid)
    for col_index in range(width):
        label = ""
        for row in reversed(grid):
            value = clean_text(row[col_index] if col_index < len(row) else "")
            if value:
                label = value
                break
        labels.append(label or fallback_header_label(col_index))

    return labels


def expanded_header_grid(
    header_rows: list[list[str]],
    header_row_numbers: list[int],
    merged_ranges: list[dict[str, Any]],
) -> list[list[str]]:
    width = max(
        [len(row) for row in header_rows] +
        [merged["max_col"] for merged in merged_ranges if merged["min_row"] in header_row_numbers or merged["max_row"] in header_row_numbers] +
        [1]
    )
    grid = [row + [""] * (width - len(row)) for row in header_rows]
    row_index_by_number = {row_number: index for index, row_number in enumerate(header_row_numbers)}

    for merged in merged_ranges:
        overlapping_rows = [
            row_number
            for row_number in header_row_numbers
            if merged["min_row"] <= row_number <= merged["max_row"]
        ]
        if not overlapping_rows:
            continue

        label = clean_text(merged["label"])
        for row_number in overlapping_rows:
            row_index = row_index_by_number[row_number]
            for column_index in range(merged["min_col"] - 1, min(merged["max_col"], width)):
                grid[row_index][column_index] = label

    return grid


def effective_header_merged_ranges(
    header_rows: list[list[str]],
    header_row_numbers: list[int],
    merged_ranges: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    row_index_by_number = {row_number: index for index, row_number in enumerate(header_row_numbers)}
    effective_ranges: list[dict[str, Any]] = []

    for merged in merged_ranges:
        overlapping_rows = [
            row_number
            for row_number in header_row_numbers
            if merged["min_row"] <= row_number <= merged["max_row"]
        ]
        if not overlapping_rows:
            continue

        adjusted = dict(merged)
        first_overlap = min(overlapping_rows)
        min_col = merged["min_col"]
        max_col = merged["max_col"]

        for candidate_row in reversed([row_number for row_number in header_row_numbers if row_number < first_overlap]):
            row_index = row_index_by_number[candidate_row]
            can_lift = True
            for column_number in range(min_col, max_col + 1):
                raw_value = clean_text(header_rows[row_index][column_number - 1] if column_number <= len(header_rows[row_index]) else "")
                if raw_value or header_cell_has_other_merge(candidate_row, column_number, merged, merged_ranges):
                    can_lift = False
                    break

            if not can_lift:
                break

            adjusted["min_row"] = candidate_row

        effective_ranges.append(adjusted)

    return effective_ranges


def header_ranges_with_vertical_singletons(
    header_rows: list[list[str]],
    header_row_numbers: list[int],
    merged_ranges: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    effective_ranges = effective_header_merged_ranges(header_rows, header_row_numbers, merged_ranges)
    grid = expanded_header_grid(header_rows, header_row_numbers, effective_ranges)

    if len(header_row_numbers) < 2 or not grid:
        return effective_ranges

    width = max(len(row) for row in grid)
    synthetic_ranges: list[dict[str, Any]] = []

    for column_index in range(width):
        non_empty = [
            (row_index, clean_text(row[column_index] if column_index < len(row) else ""))
            for row_index, row in enumerate(grid)
            if clean_text(row[column_index] if column_index < len(row) else "")
        ]

        if len(non_empty) != 1:
            continue

        row_index, label = non_empty[0]
        if row_index == 0:
            continue

        synthetic_ranges.append({
            "min_row": header_row_numbers[0],
            "max_row": header_row_numbers[-1],
            "min_col": column_index + 1,
            "max_col": column_index + 1,
            "label": label,
        })

    return effective_ranges + synthetic_ranges


def header_cell_has_other_merge(
    row_number: int,
    column_number: int,
    current_merge: dict[str, Any],
    merged_ranges: list[dict[str, Any]],
) -> bool:
    for merged in merged_ranges:
        if merged is current_merge:
            continue
        if (
            merged["min_row"] <= row_number <= merged["max_row"] and
            merged["min_col"] <= column_number <= merged["max_col"] and
            clean_text(merged["label"])
        ):
            return True

    return False


def structured_header_rows(
    header_rows: list[list[str]],
    header_row_numbers: list[int],
    merged_ranges: list[dict[str, Any]],
) -> list[list[dict[str, Any]]]:
    if not merged_ranges or not header_rows:
        return []

    effective_ranges = header_ranges_with_vertical_singletons(header_rows, header_row_numbers, merged_ranges)
    width = max(
        [len(row) for row in header_rows] +
        [merged["max_col"] for merged in effective_ranges if merged["min_row"] in header_row_numbers or merged["max_row"] in header_row_numbers] +
        [1]
    )
    row_index_by_number = {row_number: index for index, row_number in enumerate(header_row_numbers)}
    merged_by_start = {
        (merged["min_row"], merged["min_col"]): merged
        for merged in effective_ranges
    }
    covered: set[tuple[int, int]] = set()
    structured: list[list[dict[str, Any]]] = []
    grid = expanded_header_grid(header_rows, header_row_numbers, effective_ranges)

    for row_index, row_number in enumerate(header_row_numbers):
        cells: list[dict[str, Any]] = []
        for column_number in range(1, width + 1):
            if (row_number, column_number) in covered:
                continue

            merged = merged_by_start.get((row_number, column_number))
            if merged:
                row_span = len([
                    candidate
                    for candidate in header_row_numbers
                    if merged["min_row"] <= candidate <= merged["max_row"]
                ])
                col_span = min(merged["max_col"], width) - merged["min_col"] + 1
                for candidate_row in header_row_numbers:
                    if not (merged["min_row"] <= candidate_row <= merged["max_row"]):
                        continue
                    for candidate_col in range(merged["min_col"], min(merged["max_col"], width) + 1):
                        if candidate_row == row_number and candidate_col == column_number:
                            continue
                        covered.add((candidate_row, candidate_col))
                cell = {"label": clean_text(merged["label"])}
                if col_span > 1:
                    cell["colspan"] = col_span
                if row_span > 1:
                    cell["rowspan"] = row_span
                cells.append(cell)
                continue

            label = clean_text(grid[row_index][column_number - 1] if column_number <= len(grid[row_index]) else "") or fallback_header_label(column_number - 1)
            row_span = 1
            for next_row_index in range(row_index + 1, len(header_row_numbers)):
                next_row_number = header_row_numbers[next_row_index]
                next_value = clean_text(grid[next_row_index][column_number - 1] if column_number <= len(grid[next_row_index]) else "")
                if next_value:
                    break
                covered.add((next_row_number, column_number))
                row_span += 1

            cell = {"label": label}
            if row_span > 1:
                cell["rowspan"] = row_span
            cells.append(cell)
        structured.append(cells)

    return structured


def adjacent_row_values(rows: list[dict[str, Any]], index: int, offset: int) -> list[str]:
    target_index = index + offset
    if target_index >= len(rows):
        return []
    if rows[target_index]["row"] != rows[index]["row"] + offset:
        return []
    return rows[target_index]["values"]


def merge_header_rows(header_rows: list[list[str]]) -> list[str]:
    width = max(len(row) for row in header_rows)
    filled_rows = [forward_fill(row + [""] * (width - len(row))) for row in header_rows]
    labels: list[str] = []

    for index in range(width):
        parts: list[str] = []
        normalized_parts: set[str] = set()
        for row in filled_rows:
            label = clean_text(row[index])
            normalized = normalize_key(label)
            if not label or normalized in normalized_parts:
                continue
            if len(normalized) > 2 and any(normalized in existing for existing in normalized_parts):
                continue
            parts.append(label)
            normalized_parts.add(normalized)
        labels.append(" ".join(parts) or fallback_header_label(index))

    return labels


def should_merge_header_rows(primary: list[str], secondary: list[str]) -> bool:
    if not secondary:
        return False

    primary_text = " ".join(normalize_key(value) for value in primary if value)
    secondary_text = " ".join(normalize_key(value) for value in secondary if value)
    if not secondary_text:
        return False

    secondary_tokens = set(secondary_text.split())
    short_subheaders = {"h", "m", "t", "mujeres", "hombres", "total", "docentes", "administrativos"}
    has_subheaders = bool(secondary_tokens & short_subheaders)
    has_primary_groups = any(token in primary_text for token in ("periodo", "febrero", "agosto", "enero", "cantidad", "matricula"))
    has_blanks = "" in primary or "" in secondary

    return has_subheaders and (has_primary_groups or has_blanks)


def should_merge_three_header_rows(primary: list[str], secondary: list[str], tertiary: list[str]) -> bool:
    if not tertiary:
        return False

    primary_text = " ".join(normalize_key(value) for value in primary if value)
    secondary_text = " ".join(normalize_key(value) for value in secondary if value)
    tertiary_text = " ".join(normalize_key(value) for value in tertiary if value)
    if not tertiary_text:
        return False

    tertiary_tokens = set(tertiary_text.split())
    short_subheaders = {"h", "m", "t", "mujeres", "hombres", "total", "docentes", "administrativos"}
    has_subheaders = bool(tertiary_tokens & short_subheaders)
    has_primary_groups = any(
        token in primary_text
        for token in ("plantel", "programa", "periodo", "febrero", "agosto", "enero", "cantidad", "matricula", "egresados")
    )
    has_secondary_group = any(token in secondary_text for token in ("porcentaje", "titulacion", "cumplimiento", "avance"))

    return has_subheaders and has_primary_groups and (has_secondary_group or "" in primary or "" in secondary)


def forward_fill(values: list[str]) -> list[str]:
    filled: list[str] = []
    current = ""
    for value in values:
        if value:
            current = value
        filled.append(current)
    return filled


def column_from_label(label: str, seen_keys: set[str], *, key_label: str | None = None) -> dict[str, Any]:
    normalized = normalize_key(key_label or label)
    display_normalized = normalize_key(label)
    key = normalized.replace(" ", "_")[:48] or "columna"
    counter = 2
    base_key = key
    while key in seen_keys:
        key = f"{base_key}_{counter}"
        counter += 1
    seen_keys.add(key)

    semantic_text = f"{normalized} {display_normalized}"

    if any(token in semantic_text for token in READONLY_TOKENS):
        column_type = "readonly"
    elif any(token in semantic_text for token in NUMBER_TOKENS) or set(display_normalized.split()) & {"m", "h", "t"}:
        column_type = "number"
    else:
        column_type = "text"

    return {
        "key": key,
        "label": label,
        "type": column_type,
        "private": any(token in semantic_text for token in PRIVATE_TOKENS),
    }


def apply_official_calculated_columns(columns: list[dict[str, Any]]) -> None:
    """Promote known Excel formula columns to reproducible web calculations."""
    keys = {column["key"] for column in columns}

    def find_column(*tokens: str, exclude: tuple[str, ...] = ()) -> dict[str, Any] | None:
        for column in columns:
            text = normalize_key(f"{column.get('label', '')} {column.get('key', '')}")
            if all(token in text for token in tokens) and not any(token in text for token in exclude):
                return column
        return None

    attendance_column = next(
        (
            column
            for column in columns
            if "atencion" in normalize_key(f"{column['label']} {column['key']}")
        ),
        None,
    )
    if attendance_column and {
        "matricula_t",
        "cantidad_estudiantes_que_asistieron_t",
    }.issubset(keys):
        attendance_column["type"] = "calculated"
        attendance_column["calculation"] = {
            "type": "percentage",
            "numeratorKey": "cantidad_estudiantes_que_asistieron_t",
            "denominatorKey": "matricula_t",
            "decimals": 2,
        }

    egresados_mujeres = find_column("egresados", "mujeres")
    egresados_hombres = find_column("egresados", "hombres")
    egresados_total = find_column("egresados", "total")
    matricula_mujeres = find_column("matricula", "mujeres")
    matricula_hombres = find_column("matricula", "hombres")
    matricula_total = find_column("matricula", "total")
    titulacion_porcentaje = find_column("titulacion", exclude=("titulados",))

    if egresados_mujeres and egresados_hombres and egresados_total:
        egresados_total["type"] = "calculated"
        egresados_total["calculation"] = {
            "type": "sum",
            "sourceKeys": [egresados_mujeres["key"], egresados_hombres["key"]],
        }

    if matricula_mujeres and matricula_hombres and matricula_total:
        matricula_total["type"] = "calculated"
        matricula_total["calculation"] = {
            "type": "sum",
            "sourceKeys": [matricula_mujeres["key"], matricula_hombres["key"]],
        }

    if egresados_total and matricula_total and titulacion_porcentaje:
        titulacion_porcentaje["type"] = "calculated"
        titulacion_porcentaje["calculation"] = {
            "type": "percentage",
            "numeratorKey": egresados_total["key"],
            "denominatorKey": matricula_total["key"],
            "decimals": 2,
        }


def value_for_column(column: dict[str, Any], value: str) -> Any:
    if column.get("private"):
        return ""

    if is_private_cell_value(value):
        return ""

    if column["type"] == "calculated":
        return ""

    if column["type"] == "number":
        normalized = value.replace(",", "").replace("%", "").strip()
        try:
            return float(normalized) if normalized else ""
        except ValueError:
            return ""

    return value


def is_private_cell_value(value: str) -> bool:
    return bool(EMAIL_RE.search(value) or CONTACT_RE.search(value) or PHONE_NUMBER_RE.search(value))


def sanitize_header_row(values: list[str]) -> list[str]:
    labels = []
    for value in values:
        normalized = normalize_key(value)
        if "@" in value or len(value) > 140:
            continue
        if any(token in normalized.split() or token in normalized for token in HEADER_TOKENS):
            labels.append(value)
    if len(labels) >= 2:
        return labels[:14]
    return []


def first_header_rows(sheets: list[dict[str, Any]]) -> list[list[str]]:
    for sheet in sheets:
        header_rows = sheet.get("headerRows") or []
        if header_rows:
            return header_rows[:3]
    return []


def synthetic_workbook_code(
    source_path: str,
    digest: str,
    official_code: str = "",
    source_index: int = 1,
) -> str:
    stem = normalize_key(Path(source_path).stem).replace(" ", "-")[:24].strip("-")
    suffix = digest[:8].upper()
    if official_code:
        return f"{official_code}-FMT-{suffix}"
    return f"FMT-{source_index:02d}-{suffix}-{stem or 'oficial'}"


def template_from_workbook_sheets(
    code: str,
    sheets: list[dict[str, Any]],
    source_path: str,
    *,
    official_code: str = "",
    source_index: int = 1,
) -> dict[str, Any] | None:
    table_sheets = [sheet for sheet in sheets if sheet.get("table")]
    if not table_sheets:
        return None

    target_code = official_code or code
    exact_indicator_matches = [
        sheet
        for sheet in table_sheets
        if target_code in (sheet.get("_indicatorCodes") or [])
    ]
    exact_code_matches = [
        sheet
        for sheet in table_sheets
        if target_code in (sheet.get("_codes") or [])
    ]
    selection_pool = exact_indicator_matches or exact_code_matches or table_sheets

    selected = max(
        selection_pool,
        key=lambda sheet: (
            len(sheet["table"]["initialRows"]),
            len(sheet["table"]["columns"]),
            sheet.get("formulaCells", 0),
        ),
    )
    table = selected["table"]
    columns = [
        {key: value for key, value in column.items() if key != "private"}
        for column in table["columns"]
    ]
    columns, initial_rows, header_rows = normalize_extracted_table(
        columns,
        table["initialRows"],
        table.get("headerRows") or [],
    )
    source_label = unicodedata.normalize("NFC", Path(source_path).name)

    return {
        "indicatorCode": code,
        "officialCode": official_code or None,
        "indicatorName": indicator_name_from_sources(official_code or code, sheets, source_label),
        "activityLabel": activity_label_from_sources(sheets),
        "sourceLabel": source_label,
        "sourcePath": unicodedata.normalize("NFC", source_path),
        "sheetName": selected["name"],
        "groups": [],
        "headerRows": header_rows,
        "columns": columns,
        "initialRows": initial_rows,
        "showTotals": any(column["type"] in ("number", "calculated") for column in columns),
        "allowAddRows": True,
        "addRowLabel": "Agregar fila",
        "emptyRow": empty_row_for_columns(columns),
        "footerNote": "",
        "quality": [
            value
            for value in [
                "source_workbook_template",
                "private_fields_blank",
                "pending_indicator_code" if not official_code else "official_code_detected",
                "shared_official_code_split_by_source" if official_code and official_code != code else "",
            ]
            if value
        ],
    }


def indicator_name_from_sources(code: str, sheets: list[dict[str, Any]], fallback: str) -> str:
    for sheet in sheets:
        for description in sheet.get("codeDescriptions", []):
            if code not in description:
                continue
            cleaned = re.sub(r"(?i)^.*indicador\s*:?\s*", "", description).strip()
            cleaned = cleaned.replace(code, "").strip(" .:-")
            cleaned = re.sub(r"(?i)^c[oó]digo\s*", "", cleaned).strip(" .:-")
            if cleaned:
                return cleaned[:180]

    return re.sub(r"\.xlsx$", "", fallback, flags=re.I)


def activity_label_from_sources(sheets: list[dict[str, Any]]) -> str:
    for sheet in sheets:
        for activity in sheet.get("activityDescriptions", []):
            cleaned = clean_text(activity)
            if cleaned:
                return cleaned

    return ""


def empty_row_for_columns(columns: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        column["key"]: ""
        for column in columns
        if column["type"] != "calculated"
    }


def frontend_workbook_templates(
    templates: dict[str, dict[str, Any]],
    operational_names: dict[str, str],
) -> dict[str, dict[str, Any]]:
    sanitized: dict[str, dict[str, Any]] = {}
    for code in sorted(operational_names, key=normalize_key):
        template = templates.get(code)
        if not template or is_synthetic_indicator_code(code):
            continue

        sanitized[code] = {
            "indicatorCode": code,
            "officialCode": code,
            "indicatorName": operational_names[code],
            "activityLabel": "",
            "sourceLabel": "Development fallback",
            "sourcePath": "",
            "sheetName": "",
            "groups": template["groups"],
            "headerRows": template.get("headerRows") or [],
            "columns": template["columns"],
            "initialRows": template["initialRows"],
            "showTotals": template["showTotals"],
            "allowAddRows": template["allowAddRows"],
            "addRowLabel": template["addRowLabel"],
            "emptyRow": template["emptyRow"],
            "footerNote": "",
            "quality": [],
        }
    return sanitized


def classify_workbook(sheets: list[dict[str, Any]]) -> str:
    text = normalize_key(" ".join(" ".join(row) for sheet in sheets for row in sheet.get("headerRows", [])))
    if "mujer" in text and "hombre" in text and "total" in text:
        return "matrix_gender_totals"
    if "meta" in text and "avance" in text:
        return "tracking_meta_progress"
    if "correo" in text or "telefono" in text:
        return "private_roster"
    return "official_table"


def category_from_path(filename: str, source_prefix: str = "") -> str:
    parts = [part for part in unicodedata.normalize("NFC", filename).split("/") if part]
    if source_prefix == "Indicadores" and len(parts) >= 2:
        return f"Indicadores / {parts[1]}"
    if len(parts) >= 3:
        return parts[2]
    if len(parts) >= 2:
        return parts[1]
    return "General"


def generate_catalog_file(rows: list[dict[str, Any]], stats: dict[str, Any], scopes: dict[str, list[int]]) -> str:
    source_name = INDICADORES_ZIP.name
    return f"""// Generated by tools/import-official-data.py from {source_name}.
// Source binaries and row-level private data are intentionally not committed.

export type OfficialCatalogRow = {{
  sourceRow: number;
  code: string;
  sourceCode: string;
  officialCode?: string | null;
  name: string;
  responsible: string;
  contributors: string;
  activity: string;
  classification: "operational" | "template" | "template_variant" | "pending_mapping";
  visible: boolean;
  classificationReason: string[];
  dedupeKey: string;
  isDuplicate: boolean;
  duplicateOfSourceRow: number | null;
  dataQuality: string[];
}};

export const officialCatalogStats = {json_ts(stats)} as const;

export const officialIndicatorPlantelScopes: Record<string, number[]> = {json_ts(scopes)};

export const officialCatalogRows: OfficialCatalogRow[] = {json_ts(rows)};
"""


def generate_frontend_catalog_file(rows: list[dict[str, Any]], scopes: dict[str, list[int]]) -> str:
    return f"""// Generated by tools/import-official-data.py for development fallback only.
// Production catalog and assignment options are served by the backend API.

export type OfficialCatalogRow = {{
  code: string;
  name: string;
  responsible: string;
  contributors: string;
  activity: string;
}};

export const officialIndicatorPlantelScopes: Record<string, number[]> = {json_ts(scopes)};

export const officialCatalogRows: OfficialCatalogRow[] = {json_ts(rows)};
"""


def generate_data_file(
    summaries: list[dict[str, Any]],
    evidence_groups: list[dict[str, Any]],
    template_candidates: list[dict[str, Any]],
    workbook_templates: dict[str, dict[str, Any]],
    *,
    frontend: bool = False,
) -> str:
    worksheet_count = sum(len(summary["sheets"]) for summary in summaries)
    row_count = sum(sheet["nonEmptyRows"] for summary in summaries for sheet in summary["sheets"])
    nested_file_count = sum(group["fileCount"] for group in evidence_groups)
    nested_total_bytes = sum(group["totalBytes"] for group in evidence_groups)
    source_packages = [path.name for path, _scope, _prefix in official_workbook_archives()]
    summary = {
        "sourcePackage": "Development fallback" if frontend else ", ".join(source_packages) or INDICADORES_ZIP.name,
        "plantel": "Indicadores oficiales",
        "generatedAt": "2026-06-30",
        "topLevelFiles": 0 if frontend else 3 + (1 if INDICADORES_ZIP.exists() else 0),
        "nestedFiles": nested_file_count,
        "nestedTotalBytes": nested_total_bytes,
        "workbookCount": len(summaries),
        "worksheetCount": worksheet_count,
        "worksheetNonEmptyRows": row_count,
        "privacy": "Frontend fallback contains aggregate structure only; full source labels and row-level personal data are backend-private."
        if frontend
        else "Row-level personal data and evidence files remain backend-private; source binaries are not committed.",
    }
    return f"""// Generated by tools/import-official-data.py from official Drive exports.
// Source binaries and personal identifiers are intentionally not committed.

export type OfficialEvidenceGroup = {{
  category: string;
  fileCount: number;
  totalBytes: number;
  byExtension: Record<string, number>;
  sampleFileTypes: string[];
}};

export type OfficialCalculationConfig =
  | {{ type: "sum"; sourceKeys: string[] }}
  | {{
      type: "percentage";
      numeratorKey: string;
      denominatorKey: string;
      decimals?: number;
    }}
  | {{
      type: "formula";
      expression: string;
      decimals?: number;
    }};

export type OfficialWorkbookSheetSummary = {{
  name: string;
  nonEmptyRows: number;
  columnsObserved: number;
  sampleHeaders: string[];
  headerRows: string[][];
  numericCells: number;
  textCells: number;
  formulaCells: number;
  table?: {{
    headerRow: number;
    headerRows?: Array<Array<{{
      label: string;
      colspan?: number;
      rowspan?: number;
    }}>>;
    columns: Array<{{
      key: string;
      label: string;
      type: "readonly" | "number" | "text" | "calculated";
      calculation?: OfficialCalculationConfig;
      private?: boolean;
    }}>;
    initialRows: Array<Record<string, unknown>>;
  }} | null;
  codeDescriptions: string[];
  activityDescriptions: string[];
}};

export type OfficialWorkbookSummary = {{
  id: string;
  sourceLabel: string;
  sourcePathHash: string;
  category: string;
  sizeBytes: number;
  detectedIndicatorCodes: string[];
  detectedReferenceCodes: string[];
  detectedPlanteles: string[];
  formulaCells: number;
  sheets: OfficialWorkbookSheetSummary[];
  privacy: string;
}};

export type OfficialTemplateCandidate = {{
  indicatorCode: string;
  officialCode?: string | null;
  sourcePath: string;
  planteles: string[];
  classification: string;
  confidence: string;
  headerRows: string[][];
}};

export type OfficialWorkbookTemplate = {{
  indicatorCode: string;
  officialCode?: string | null;
  indicatorName: string;
  sourceLabel: string;
  activityLabel?: string;
  sourcePath: string;
  sheetName: string;
  groups: Array<{{ label: string; colspan: number }}>;
  headerRows?: Array<Array<{{
    label: string;
    colspan?: number;
    rowspan?: number;
  }}>>;
  columns: Array<{{
    key: string;
    label: string;
    type: "readonly" | "number" | "text" | "calculated";
    calculation?: OfficialCalculationConfig;
  }}>;
  initialRows: Array<Record<string, unknown>>;
  showTotals: boolean;
  allowAddRows: boolean;
  addRowLabel: string;
  emptyRow: Record<string, unknown>;
  footerNote: string;
  quality: string[];
}};

export type OfficialDataSummary = {{
  sourcePackage: string;
  plantel: string;
  generatedAt: string;
  topLevelFiles: number;
  nestedFiles: number;
  nestedTotalBytes: number;
  workbookCount: number;
  worksheetCount: number;
  worksheetNonEmptyRows: number;
  privacy: string;
}};

export const officialDataSummary: OfficialDataSummary = {json_ts(summary)};

export const officialEvidenceGroups: OfficialEvidenceGroup[] = {json_ts(evidence_groups)};

export const officialTemplateCandidates: OfficialTemplateCandidate[] = {json_ts(template_candidates)};

export const officialWorkbookTemplates: Record<string, OfficialWorkbookTemplate> = {json_ts(workbook_templates)};

export const officialWorkbookSummaries: OfficialWorkbookSummary[] = {json_ts(summaries)};
"""


def main() -> None:
    summaries, detected_scopes, evidence_groups, template_candidates, workbook_templates = workbook_summaries()
    workbook_templates = canonicalize_workbook_templates(workbook_templates)
    rows = workbook_catalog_rows(workbook_templates)
    stats = catalog_stats(rows)
    scopes = catalog_plantel_scopes(rows, detected_scopes)

    public_catalog_rows = frontend_catalog_rows(rows)
    operational_names = {
        row["code"]: row["name"]
        for row in public_catalog_rows
    }
    public_scopes = {
        code: plantel_ids
        for code, plantel_ids in scopes.items()
        if code in operational_names
    }
    backend_catalog_text = generate_catalog_file(rows, stats, scopes)
    frontend_catalog_text = generate_frontend_catalog_file(public_catalog_rows, public_scopes)
    backend_data_text = generate_data_file(summaries, evidence_groups, template_candidates, workbook_templates)
    frontend_data_text = generate_data_file(
        [],
        [],
        [],
        frontend_workbook_templates(workbook_templates, operational_names),
        frontend=True,
    )
    classification_matrix_text = generate_classification_matrix(rows, stats)

    frontend_only = os.environ.get("ADPEAK_FRONTEND_ONLY") == "1"

    if not frontend_only:
        write_text(BACKEND_CATALOG_TARGET, backend_catalog_text)
        write_text(BACKEND_DATA_TARGET, backend_data_text)
        write_text(CLASSIFICATION_MATRIX_TARGET, classification_matrix_text)

    write_text(FRONTEND_CATALOG_TARGET, frontend_catalog_text)
    write_text(FRONTEND_DATA_TARGET, frontend_data_text)

    print(json.dumps({
        "catalog": stats,
        "workbooks": len(summaries),
        "scopedIndicators": len(scopes),
        "workbookTemplates": len(workbook_templates),
        "frontendOnly": frontend_only,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
